"""Document loaders for PDF, DOCX, TXT and Markdown that keep page/section metadata."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class PageBlock:
    """A contiguous block of text with its location in the source document."""

    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class LoadedDocument:
    doc_id: str
    source: str
    blocks: list[PageBlock] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)


def _doc_id(path: Path, data: bytes) -> str:
    return hashlib.sha1(path.name.encode() + data).hexdigest()[:16]


_HEADING_RE = re.compile(r"^(?:\d+(?:\.\d+)*\.?\s+)?[A-Z][A-Za-z0-9 ,&/\-()]{2,80}$")


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 80 or line.endswith((".", ",", ";", ":")):
        return False
    return bool(_HEADING_RE.match(line)) and len(line.split()) <= 10


def load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader

    data = path.read_bytes()
    reader = PdfReader(str(path))
    doc = LoadedDocument(doc_id=_doc_id(path, data), source=path.name)
    section: str | None = None
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        buf: list[str] = []
        for line in text.splitlines():
            if _looks_like_heading(line):
                if buf:
                    doc.blocks.append(PageBlock("\n".join(buf).strip(), page=i, section=section))
                    buf = []
                section = line.strip()
            buf.append(line)
        if buf and "\n".join(buf).strip():
            doc.blocks.append(PageBlock("\n".join(buf).strip(), page=i, section=section))
    return doc


def load_docx(path: Path) -> LoadedDocument:
    import docx

    data = path.read_bytes()
    d = docx.Document(str(path))
    doc = LoadedDocument(doc_id=_doc_id(path, data), source=path.name)
    section: str | None = None
    buf: list[str] = []
    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            if buf:
                doc.blocks.append(PageBlock("\n".join(buf), section=section))
                buf = []
            section = text
        buf.append(text)
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            buf.append("\n".join(rows))
    if buf:
        doc.blocks.append(PageBlock("\n".join(buf), section=section))
    return doc


def load_text(path: Path) -> LoadedDocument:
    data = path.read_bytes()
    doc = LoadedDocument(doc_id=_doc_id(path, data), source=path.name)
    section: str | None = None
    buf: list[str] = []
    for line in data.decode("utf-8", errors="ignore").splitlines():
        if line.startswith("#"):
            if "\n".join(buf).strip():
                doc.blocks.append(PageBlock("\n".join(buf).strip(), section=section))
            buf = []
            section = line.lstrip("#").strip()
        buf.append(line)
    if "\n".join(buf).strip():
        doc.blocks.append(PageBlock("\n".join(buf).strip(), section=section))
    return doc


def load_document(path: str | Path) -> LoadedDocument:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in {".txt", ".md"}:
        return load_text(path)
    raise ValueError(f"Unsupported file type: {ext}")
