"""Generate the demo corpus (PDF/DOCX/MD), the SQLite sales DB and the golden eval set.

    python scripts/generate_sample_data.py
"""

from __future__ import annotations

import json
import random
import sqlite3
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from corpus import DOCS, PRODUCTS, STORES  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "data" / "sample_docs"
DB_PATH = ROOT / "data" / "sample.db"
GOLDEN = ROOT / "eval" / "golden_set.jsonl"
EXT = {"pdf": ".pdf", "docx": ".docx", "md": ".md"}


def write_pdf(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for i, (heading, body) in enumerate(sections):
        if i and i % 2 == 0:
            story.append(PageBreak())  # multi-page docs so page citations are meaningful
        story += [Paragraph(heading, styles["Heading2"]), Paragraph(body, styles["BodyText"]), Spacer(1, 10)]
    SimpleDocTemplate(str(path), pagesize=letter, title=title).build(story)


def write_docx(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    import docx

    d = docx.Document()
    d.add_heading(title, 0)
    for heading, body in sections:
        d.add_heading(heading, 1)
        d.add_paragraph(body)
    d.save(str(path))


def write_md(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    lines = [f"# {title}", ""]
    for heading, body in sections:
        lines += [f"## {heading}", "", body, ""]
    path.write_text("\n".join(lines))


def build_docs() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    writers = {"pdf": write_pdf, "docx": write_docx, "md": write_md}
    for d in DOCS:
        writers[d["fmt"]](DOCS_DIR / f"{d['file']}{EXT[d['fmt']]}", d["title"], d["sections"])


def build_db() -> None:
    DB_PATH.unlink(missing_ok=True)
    rng = random.Random(42)
    conn = sqlite3.connect(DB_PATH)
    conn.executescript("""
        CREATE TABLE products (sku TEXT PRIMARY KEY, name TEXT NOT NULL, category TEXT NOT NULL, unit_price REAL NOT NULL);
        CREATE TABLE stores (store_id INTEGER PRIMARY KEY, name TEXT NOT NULL, city TEXT NOT NULL, region TEXT NOT NULL);
        CREATE TABLE sales (id INTEGER PRIMARY KEY, sale_date TEXT NOT NULL, store_id INTEGER NOT NULL REFERENCES stores(store_id),
                            sku TEXT NOT NULL REFERENCES products(sku), units INTEGER NOT NULL, revenue REAL NOT NULL);
        CREATE TABLE inventory (store_id INTEGER NOT NULL, sku TEXT NOT NULL, on_hand INTEGER NOT NULL, reorder_point INTEGER NOT NULL,
                                PRIMARY KEY (store_id, sku));
        CREATE TABLE recalls (sku TEXT PRIMARY KEY REFERENCES products(sku), recall_class TEXT NOT NULL, reason TEXT NOT NULL, announced_on TEXT NOT NULL);
    """)
    conn.executemany("INSERT INTO products VALUES (?,?,?,?)", PRODUCTS)
    conn.executemany("INSERT INTO stores VALUES (?,?,?,?)", STORES)
    conn.executemany("INSERT INTO recalls VALUES (?,?,?,?)", [
        ("HF-1042", "Class I", "possible Listeria", "2026-06-03"),
        ("HF-2210", "Class II", "undeclared peanut", "2026-05-21"),
        ("HF-3307", "Class II", "possible Hepatitis A", "2026-06-10"),
    ])
    start = date(2026, 1, 1)
    rows = []
    for day in range(181):  # Jan 1 - Jun 30 2026
        d = start + timedelta(days=day)
        for store_id, *_ in STORES:
            for sku, _, _, price in PRODUCTS:
                if rng.random() < 0.55:
                    units = rng.randint(1, 40)
                    rows.append((d.isoformat(), store_id, sku, units, round(units * price, 2)))
    conn.executemany("INSERT INTO sales (sale_date, store_id, sku, units, revenue) VALUES (?,?,?,?,?)", rows)
    conn.executemany("INSERT INTO inventory VALUES (?,?,?,?)", [
        (s[0], p[0], rng.randint(0, 200), 40) for s in STORES for p in PRODUCTS
    ])
    conn.commit()
    conn.close()


def build_golden() -> None:
    GOLDEN.parent.mkdir(parents=True, exist_ok=True)
    with GOLDEN.open("w") as f:
        for d in DOCS:
            for q, a, section in d["qa"]:
                f.write(json.dumps({"question": q, "answer": a, "source": f"{d['file']}{EXT[d['fmt']]}",
                                    "section": section}) + "\n")


if __name__ == "__main__":
    build_docs()
    build_db()
    build_golden()
    n_q = sum(len(d["qa"]) for d in DOCS)
    print(f"Wrote {len(DOCS)} documents to {DOCS_DIR}, sales DB to {DB_PATH}, {n_q} eval questions to {GOLDEN}")
